"""Laporan dihitung saat diminta dari tabel lain; yang disimpan hanya catatan unduhan (report_downloads)."""
import io
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, literal_column, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.alerts.models import LEVELS, Alert, AlertRule
from app.alerts.rules import describe
from app.auth import audit
from app.auth.models import User
from app.core.templating import fmt_bytes, localtime
from app.folders.service import FolderTree
from app.hosts.models import Host, HostDisk, MetricHourly, RemoteSession, Screenshot
from app.hosts.service import APPROVED
from app.reports.models import ReportDownload

REPORTS = {
    "inventory": "Inventaris hardware",
    "availability": "Status & ketersediaan PC",
    "alerts": "Rekap alert",
    "renewal": "PC perlu peremajaan",
    "remote": "Aktivitas remote & screenshot",
}
SENSITIVE = {"remote"}               # hanya Admin (hak akses report_sensitive)
RANGED = {"availability", "alerts", "remote"}
MAX_ROWS = 5000
RENEWAL_MIN_RAM_MB = 7900            # < 8 GB
RENEWAL_MIN_DISK = 200 * 1000**3


@dataclass
class Report:
    key: str
    title: str
    columns: list[str]
    rows: list[list] = field(default_factory=list)
    subtitle: str = ""


@dataclass
class Params:
    folder_ids: list[int] | None      # None = semua folder
    folder_label: str
    start: datetime | None = None     # UTC
    end: datetime | None = None
    period_label: str = ""

    def as_json(self) -> dict:
        return {"folder_ids": self.folder_ids, "from": self.start.isoformat() if self.start else None,
                "to": self.end.isoformat() if self.end else None}


def _host_scope(stmt, p: Params):
    stmt = stmt.where(Host.status.in_(APPROVED))
    return stmt.where(Host.folder_id.in_(p.folder_ids)) if p.folder_ids is not None else stmt


def _gb(mb: int | None) -> str:
    return f"{round(mb / 1024)} GB" if mb else "—"


async def _disks(db: AsyncSession, host_ids: list[int]) -> dict[int, list[HostDisk]]:
    out: dict[int, list[HostDisk]] = {}
    if host_ids:
        for d in await db.scalars(select(HostDisk).where(HostDisk.host_id.in_(host_ids)).order_by(HostDisk.mount)):
            out.setdefault(d.host_id, []).append(d)
    return out


def _disk_text(disks: list[HostDisk]) -> str:
    return " · ".join(f"{d.mount} {fmt_bytes(d.total_bytes)}{' ' + d.media if d.media else ''}" for d in disks) or "—"


async def inventory(db: AsyncSession, tree: FolderTree, p: Params) -> Report:
    hosts = list(await db.scalars(_host_scope(select(Host), p).order_by(Host.hostname).limit(MAX_ROWS)))
    disks = await _disks(db, [h.id for h in hosts])
    r = Report("inventory", f"Inventaris Hardware · {p.folder_label}",
               ["Hostname", "Folder", "Merek / model", "CPU", "RAM", "Disk", "OS", "Serial", "No. inventaris"])
    for h in hosts:
        r.rows.append([h.name, tree.path(h.folder_id), " ".join(x for x in (h.manufacturer, h.model) if x) or "—",
                       h.cpu_model or "—", _gb(h.ram_total_mb), _disk_text(disks.get(h.id, [])),
                       " ".join(x for x in (h.os_name, h.os_build) if x) or "—", h.serial_number or "—",
                       h.inventory_no or "—"])
    return r


async def availability(db: AsyncSession, tree: FolderTree, p: Params) -> Report:
    # Penyebut = menit sejak PC terdaftar (bukan seluruh periode) sampai akhir periode atau sekarang
    minutes = func.timestampdiff(literal_column("MINUTE"), func.greatest(p.start, Host.created_at),
                                 func.least(p.end, func.now()))
    online = (select(MetricHourly.host_id, func.sum(MetricHourly.online_min).label("online"))
              .where(MetricHourly.hour >= p.start, MetricHourly.hour < p.end)
              .group_by(MetricHourly.host_id).subquery())
    stmt = _host_scope(
        select(Host, func.coalesce(online.c.online, 0), minutes).outerjoin(online, online.c.host_id == Host.id), p)
    rows = (await db.execute(stmt.order_by(Host.hostname).limit(MAX_ROWS))).all()
    r = Report("availability", f"Status & Ketersediaan PC · {p.folder_label}",
               ["Hostname", "Folder", "Ketersediaan", "Jam online", "Status", "Last seen"], subtitle=p.period_label)
    rows = sorted(rows, key=lambda row: (float(row[1]) / row[2]) if row[2] else 0)
    for h, online_min, total in rows:
        pct = float(online_min) / total * 100 if total else 0
        r.rows.append([h.name, tree.path(h.folder_id), f"{pct:.1f}%".replace(".", ","),
                       f"{float(online_min) / 60:.0f}", h.status, localtime(h.last_seen_at)])
    return r


async def alerts(db: AsyncSession, tree: FolderTree, p: Params) -> Report:
    stmt = (select(AlertRule, Alert.level, func.count(), func.count(func.distinct(Alert.host_id)),
                   func.avg(func.timestampdiff(literal_column("SECOND"), Alert.started_at,
                                               func.coalesce(Alert.resolved_at, func.now()))))
            .join(AlertRule, AlertRule.id == Alert.rule_id).join(Host, Host.id == Alert.host_id)
            .where(Alert.started_at >= p.start, Alert.started_at < p.end)
            .group_by(AlertRule.id, Alert.level).order_by(func.count().desc()))
    if p.folder_ids is not None:
        stmt = stmt.where(Host.folder_id.in_(p.folder_ids))
    r = Report("alerts", f"Rekap Alert · {p.folder_label}",
               ["Aturan", "Kondisi", "Level", "Jumlah alert", "PC terdampak", "Rata-rata durasi (menit)"],
               subtitle=p.period_label)
    for rule, level, count, hosts, avg in (await db.execute(stmt)).all():
        r.rows.append([rule.name, describe(rule), LEVELS[level], count, hosts, round(float(avg or 0) / 60)])
    return r


async def renewal(db: AsyncSession, tree: FolderTree, p: Params) -> Report:
    hosts = list(await db.scalars(_host_scope(select(Host), p).order_by(Host.hostname).limit(MAX_ROWS)))
    disks = await _disks(db, [h.id for h in hosts])
    r = Report("renewal", f"PC Perlu Peremajaan · {p.folder_label}",
               ["Hostname", "Folder", "Merek / model", "RAM", "Disk", "OS", "Alasan"],
               subtitle="Kriteria: RAM < 8 GB, disk sistem HDD atau < 200 GB, Windows 7/8")
    for h in hosts:
        reasons = []
        if h.ram_total_mb and h.ram_total_mb < RENEWAL_MIN_RAM_MB:
            reasons.append(f"RAM {_gb(h.ram_total_mb)}")
        system = next((d for d in disks.get(h.id, []) if d.mount.upper().startswith("C")), None)
        if system and (system.media or "").upper() == "HDD":
            reasons.append("Disk sistem HDD")
        if system and system.total_bytes < RENEWAL_MIN_DISK:
            reasons.append(f"Disk sistem {fmt_bytes(system.total_bytes)}")
        if h.os_name and any(v in h.os_name for v in ("Windows 7", "Windows 8")):
            reasons.append(h.os_name)
        if reasons:
            r.rows.append([h.name, tree.path(h.folder_id), " ".join(x for x in (h.manufacturer, h.model) if x) or "—",
                           _gb(h.ram_total_mb), _disk_text(disks.get(h.id, [])), h.os_name or "—", ", ".join(reasons)])
    return r


async def remote(db: AsyncSession, tree: FolderTree, p: Params) -> Report:
    r = Report("remote", f"Aktivitas Remote & Screenshot · {p.folder_label}",
               ["Waktu", "User", "Jenis", "PC", "Detail", "Alasan"], subtitle=p.period_label)
    shots = select(Screenshot.requested_at, User.username, Host.hostname, Screenshot.status, Screenshot.reason) \
        .join(User, User.id == Screenshot.requested_by).join(Host, Host.id == Screenshot.host_id) \
        .where(Screenshot.requested_at >= p.start, Screenshot.requested_at < p.end)
    sessions = select(RemoteSession.started_at, User.username, Host.hostname, RemoteSession.mode,
                      RemoteSession.reason, RemoteSession.ended_at) \
        .join(User, User.id == RemoteSession.user_id).join(Host, Host.id == RemoteSession.host_id) \
        .where(RemoteSession.started_at >= p.start, RemoteSession.started_at < p.end)
    if p.folder_ids is not None:
        shots = shots.where(Host.folder_id.in_(p.folder_ids))
        sessions = sessions.where(Host.folder_id.in_(p.folder_ids))
    items = [(ts, user, "Screenshot", host, status, reason) for ts, user, host, status, reason
             in (await db.execute(shots.limit(MAX_ROWS))).all()]
    for ts, user, host, mode, reason, ended in (await db.execute(sessions.limit(MAX_ROWS))).all():
        dur = f" · {round((ended - ts).total_seconds() / 60)} mnt" if ended else " · berjalan"
        items.append((ts, user, "Remote", host, ("Kontrol penuh" if mode == "control" else "Lihat saja") + dur, reason))
    for ts, user, kind, host, detail, reason in sorted(items, key=lambda i: i[0], reverse=True):
        r.rows.append([localtime(ts), user, kind, host, detail, reason])
    return r


BUILDERS = {"inventory": inventory, "availability": availability, "alerts": alerts, "renewal": renewal,
            "remote": remote}


async def build(db: AsyncSession, key: str, tree: FolderTree, p: Params) -> Report:
    return await BUILDERS[key](db, tree, p)


async def record_download(db: AsyncSession, key: str, fmt: str, p: Params, actor: audit.Actor) -> None:
    async with db.begin():
        db.add(ReportDownload(user_id=actor.user_id, report=key, format=fmt, params=p.as_json()))
        await audit.record(db, actor, "report", "Unduh laporan", sensitive=key in SENSITIVE,
                           target_label=REPORTS[key][:80], detail=f"{fmt.upper()} · {p.folder_label}"
                           + (f" · {p.period_label}" if p.period_label else ""))


def to_xlsx(report: Report, org_name: str, printed: str) -> bytes:
    """CPU-bound kecil (ratusan baris); panggil lewat run_in_threadpool."""
    wb = Workbook()
    ws = wb.active
    ws.title = report.key
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"{org_name} · dicetak {printed}" + (f" · {report.subtitle}" if report.subtitle else "")])
    ws.append([])
    ws.append(report.columns)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for row in report.rows:
        ws.append(row)
    for col in ws.iter_cols(min_row=4):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

